import mysql.connector
from datetime import datetime
from discord_webhook import DiscordWebhook, DiscordEmbed
from dotenv import load_dotenv
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import os


AIRFLOW_HOME=os.environ.get("AIRFLOW_HOME")
load_dotenv(os.path.join(
    AIRFLOW_HOME,
    "..",
    ".env"
))
WEBHOOK_URL = os.environ.get("WEBHOOK_URL")

processing_date = datetime.today().date().strftime(f"%Y-%m-%d")


def get_jobs(processing_date=None):
    conn = mysql.connector.connect(
        host="mysql",
        port=3306,
        user="mysql",
        password="mysql",
        database="jobs_db"
    )

    query = """
        SELECT
            job_title,
            company,
            address,
            salary,
            link_description
        FROM (
            SELECT
                j.job_title,
                j.company,
                j.address,
                j.salary,
                j.link_description,
                ROW_NUMBER() OVER (
                    PARTITION BY j.link_description
                    ORDER BY j.created_date DESC
                ) AS rn
            FROM jobs AS j
            WHERE FIND_IN_SET(
                %s,
                REPLACE(j.tag, ', ', ',')
            ) > 0
            AND DATE(j.created_date) = %s
        ) AS t
        WHERE rn = 1;
    """


    try:
        cursor = conn.cursor(dictionary=True)

        cursor.execute(
            query,
            ("Data Engineer", processing_date)
        )

        jobs = cursor.fetchall()

        cursor.close()

        return jobs

    finally:
        conn.close()

def format_job(job, index):

    title = job.get("job_title") or "Không có tiêu đề"
    company = job.get("company") or "Không có thông tin"
    address = job.get("address") or "Không có thông tin"
    salary = job.get("salary")
    url = job.get("link_description")

    if url:
        link = f"[🔗 Xem job]({url})"
    else:
        link = "🔗 Không có link"

    return (
        f"**{index}. {title}**\n"
        f"🏢 {company}\n"
        f"📍 {address}\n"
        f"💰 {salary}\n"
        f"{link}"
    )


def send_to_discord(jobs, processing_date):

    if not jobs:
        print("Không có Data Engineer job.")
        return

    output = BytesIO()

    wb = Workbook()
    ws = wb.active
    ws.title = "Data Engineer Jobs"

    headers = [
        "STT",
        "Job Title",
        "Company",
        "Address",
        "Salary",
        "Link"
    ]

    ws.append(headers)

    # Header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Data
    for index, job in enumerate(jobs, start=1):

        ws.append([
            index,
            job.get("job_title") or "",
            job.get("company") or "",
            job.get("address") or "",
            job.get("salary") or "Thỏa thuận",
            job.get("link_description") or ""
        ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 45
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 60

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # Hyperlink
    for row in range(2, ws.max_row + 1):

        link_cell = ws[f"F{row}"]

        if link_cell.value:
            link_cell.hyperlink = link_cell.value
            link_cell.style = "Hyperlink"

    # Ghi workbook vào RAM
    wb.save(output)

    # Đưa con trỏ về đầu BytesIO
    output.seek(0)

    webhook = DiscordWebhook(
        url=WEBHOOK_URL,
        username="Airflow"
    )

    embed = DiscordEmbed(
        title=f"🚀 Data Engineer Jobs — {processing_date}",
        description=(
            f"📊 **{len(jobs)} jobs mới**\n\n"
            f"📎 File Excel chứa toàn bộ "
            f"**{len(jobs)} jobs** được đính kèm."
        )
    )

    embed.add_embed_field(
        name="📅 Ngày",
        value=str(processing_date),
        inline=True
    )

    embed.add_embed_field(
        name="💼 Category",
        value="Data Engineer",
        inline=True
    )

    embed.add_embed_field(
        name="📊 Số lượng",
        value=str(len(jobs)),
        inline=True
    )

    embed.set_footer(
        text="Data Engineer Airflow"
    )

    webhook.add_embed(embed)

    file_name = (
        f"data_engineer_jobs_{processing_date}.xlsx"
    )

    webhook.add_file(
        file=output.getvalue(),
        filename=file_name
    )

    response = webhook.execute()

    if response.status_code >= 300:
        raise RuntimeError(
            f"Discord error: "
            f"{response.status_code} "
            f"{response.text}"
        )

    print(
        f"Đã gửi {len(jobs)} Data Engineer jobs "
        f"và file Excel vào Discord."
    )


def main():


    print(
        f"Đang lấy Data Engineer jobs "
        f"ngày {processing_date}..."
    )

    jobs = get_jobs(processing_date)

    print(
        f"Tìm thấy {len(jobs)} jobs."
    )

    send_to_discord(
        jobs,
        processing_date
    )


if __name__ == "__main__":
    main()